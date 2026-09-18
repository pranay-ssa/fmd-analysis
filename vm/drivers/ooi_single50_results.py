import json
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# Load all metrics
archs = ['resnet18', 'resnet50', 'resnet50_inception', 'resnet50_se']
results = []
for arch in archs:
    with open(f'ooi_single_model_runs/{arch}/epochs_50/metrics.json') as f:
        results.append(json.load(f))

# Class abbreviations matching teammate's format
class_abbr = {
    'Bubble': 'B',
    'Bubble Cluster': 'BC',
    'Bubble Irregular': 'BI',
    'Bubble On 123': 'B123',
    'Bubble On Edge': 'BOE',
    'Extraneous Polymer': 'EP',
    'Fiber': 'Fl',
    'Foreign Matter': 'FM',
    'HEMA Fragment': 'HF',
    'Wet Package': 'WP',
}
class_order = list(class_abbr.keys())

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'OOI Single-Scale Results'

# Styles
header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font_white = Font(bold=True, size=11, color='FFFFFF')
yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

# Headers
headers = ['Dataset', 'Architecture', 'Overall Accuracy (%)',
           'B', 'BC', 'BI', 'B123', 'BOE', 'EP', 'Fl', 'FM', 'HF', 'WP',
           'Training time (s)', 'Inference Time per Image(ms)']

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=2, column=col, value=h)
    cell.font = header_font_white
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', wrap_text=True)
    cell.border = thin_border

# Row 1: merge top header
ws.merge_cells('A1:C1')
ws.cell(row=1, column=1, value='OOI Classification - Single-Scale CNN Results (50 Epochs, Batch 8, Seed 42)')
ws.cell(row=1, column=1).font = Font(bold=True, size=13)
ws.cell(row=1, column=1).alignment = Alignment(horizontal='center')

# Recall sub-header
ws.merge_cells('D1:M1')
ws.cell(row=1, column=4).value = 'Recall (%)'
ws.cell(row=1, column=4).font = Font(bold=True, size=11)
ws.cell(row=1, column=4).alignment = Alignment(horizontal='center')

# Data rows
dataset = '/home/pranayp/OoI_dataset_single'
arch_display = {
    'resnet18': 'ResNet18 (50 Epochs)',
    'resnet50': 'ResNet50 (50 Epochs)',
    'resnet50_inception': 'ResNet50+Inception (50 Epochs)',
    'resnet50_se': 'ResNet50+SE (50 Epochs)',
}

for row_idx, r in enumerate(results, 3):
    arch_key = r['arch']
    ws.cell(row=row_idx, column=1, value=dataset).border = thin_border
    ws.cell(row=row_idx, column=2, value=arch_display[arch_key]).border = thin_border
    ws.cell(row=row_idx, column=3, value=round(r['overall_accuracy'] * 100, 2)).border = thin_border

    # Recall per class
    for col_offset, cls in enumerate(class_order):
        recall_val = round(r['per_class'][cls]['recall'] * 100, 2)
        cell = ws.cell(row=row_idx, column=4 + col_offset, value=recall_val)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')

    # Training time and inference
    ws.cell(row=row_idx, column=14, value=round(r['training_time_sec'], 2)).border = thin_border
    ws.cell(row=row_idx, column=15, value=round(r['inference_per_image_ms'], 3)).border = thin_border

# Column widths
col_widths = [38, 30, 18, 8, 8, 8, 8, 8, 8, 8, 8, 8, 8, 18, 28]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# Save
out = 'ooi_single50_results.xlsx'
wb.save(out)
print(f'Saved {out}')

#!/usr/bin/env python3
"""Build the two candidate label matrices and one workbook for the lead to pick from.

Outputs (next to this file):
  labels_178_multitag.csv    rows carrying exactly 2 LP tags
  labels_3377_lp_tagged.csv  rows carrying at least 1 LP tag
  label_matrix_options.xlsx   both sheets + decision sheet + counts

Also prints a check of the 67 / 111 / 178 split quoted from the lead.
"""
import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
sys.path.insert(0, str(HERE))

from build_labels import LP_TAGS, parse_export, DEFAULT_XML  # noqa: E402

PRETTY = {
    "P1_COC": "Cavity_Off_Center",
    "P1_FL": "Folded_Lens",
    "P1_ML": "Multiple_Lenses",
    "P1_MPP": "Missing_Primary_Package",
    "P1_NL": "Missing_Lens",
}
LABEL_COLS = [PRETTY[t] for t in LP_TAGS]
EXTRA = ["n_lp_tags", "also_has_other_tag", "source_task"]
OUT_HEADER = ["image_path"] + LABEL_COLS + EXTRA

CSV_A = HERE / "labels_178_multitag.csv"
CSV_B = HERE / "labels_3377_lp_tagged.csv"
XLSX = HERE / "label_matrix_options.xlsx"


def load():
    """Return rows built from labels.csv plus a non-LP tag flag from the XML."""
    records, _ = parse_export(Path(DEFAULT_XML))
    other_tags = {r["image_path"]: len(r["tag_set"] - set(LP_TAGS)) for r in records}

    rows = []
    with open(HERE / "labels.csv", newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            labels = [int(r[t]) for t in LP_TAGS]
            n_lp = sum(labels)
            if n_lp == 0:
                continue
            rows.append({
                "image_path": r["image_path"],
                "labels": labels,
                "n_lp_tags": n_lp,
                "also_has_other_tag": int(other_tags.get(r["image_path"], 0) > 0),
                "source_task": r["task_name"],
            })
    rows.sort(key=lambda x: (-x["n_lp_tags"], x["image_path"]))
    return rows


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(OUT_HEADER)
        for r in rows:
            w.writerow([r["image_path"]] + r["labels"] + [
                r["n_lp_tags"], r["also_has_other_tag"], r["source_task"]])
    return path


# ---- workbook styling ------------------------------------------------------
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
BOLD = Font(bold=True)
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
OPT_A_FILL = PatternFill("solid", fgColor="FFF2CC")
OPT_B_FILL = PatternFill("solid", fgColor="E2EFDA")


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[row].height = 22


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def data_sheet(wb, title, rows, tab_color):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    ws.append(OUT_HEADER)
    for r in rows:
        ws.append([r["image_path"]] + r["labels"] + [
            r["n_lp_tags"], r["also_has_other_tag"], r["source_task"]])
    style_header(ws, len(OUT_HEADER))
    set_widths(ws, [78, 18, 13, 17, 25, 13, 12, 20, 14])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(OUT_HEADER)), ws.max_row)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=8):
        for cell in row:
            cell.alignment = TOP
    return ws


def decision_sheet(wb, a, b):
    ws = wb.active
    ws.title = "Pick one"
    ws.sheet_properties.tabColor = "1F3864"
    set_widths(ws, [34, 34, 34, 34])

    r = 1
    ws.cell(r, 1, "Multi-label CNN label matrix: which rows should train?").font = TITLE_FONT
    r += 1
    ws.cell(r, 1, "Lens presentation (LP) tags, ResNet50, prepared 2026-09-24").font = Font(
        italic=True, size=10, color="595959")
    r += 2

    ws.cell(r, 1, "Both data sheets carry the same 5 label columns. Only the row set differs.").font = BOLD
    r += 2

    heads = ["", "Option A (sheet: Option A - 178 rows)", "Option B (sheet: Option B - 3377 rows)"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(r, i, h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = WRAP
    ws.row_dimensions[r].height = 30
    r += 1

    both = Counter(tuple(x["labels"]) for x in b)
    a_pat = len(set(tuple(x["labels"]) for x in a))
    pos_a = [sum(x["labels"][i] for x in a) for i in range(5)]
    pos_b = [sum(x["labels"][i] for x in b) for i in range(5)]

    table = [
        ("Rows", "%d images" % len(a), "%d images" % len(b)),
        ("Row definition", "images carrying exactly 2 LP tags",
         "images carrying at least 1 LP tag"),
        ("Label combinations present", "%d (Folded Lens + Multiple Lenses only)" % a_pat,
         "%d" % len(both)),
        ("Folded Lens positives / negatives",
         "%d / %d" % (pos_a[1], len(a) - pos_a[1]),
         "%d / %d" % (pos_b[1], len(b) - pos_b[1])),
        ("Multiple Lenses positives / negatives",
         "%d / %d" % (pos_a[2], len(a) - pos_a[2]),
         "%d / %d" % (pos_b[2], len(b) - pos_b[2])),
        ("Cavity Off Center positives / negatives",
         "%d / %d" % (pos_a[0], len(a) - pos_a[0]),
         "%d / %d" % (pos_b[0], len(b) - pos_b[0])),
        ("Missing Primary Package positives / negatives",
         "%d / %d" % (pos_a[3], len(a) - pos_a[3]),
         "%d / %d" % (pos_b[3], len(b) - pos_b[3])),
        ("Missing Lens positives / negatives",
         "%d / %d" % (pos_a[4], len(a) - pos_a[4]),
         "%d / %d" % (pos_b[4], len(b) - pos_b[4])),
    ]
    for label, va, vb in table:
        ws.cell(r, 1, label).font = BOLD
        ws.cell(r, 1).alignment = WRAP
        ws.cell(r, 1).border = BOX
        c2 = ws.cell(r, 2, va)
        c2.alignment = WRAP
        c2.border = BOX
        c2.fill = OPT_A_FILL
        c3 = ws.cell(r, 3, vb)
        c3.alignment = WRAP
        c3.border = BOX
        c3.fill = OPT_B_FILL
        r += 1

    r += 1
    ws.cell(r, 1, "What each option supports").font = TITLE_FONT
    r += 1

    notes = [
        ("Option A", "Every row in Option A carries the same label vector: Folded Lens 1, "
                     "Multiple Lenses 1, the other three columns 0. Three columns never see a "
                     "positive example and two columns never see a negative example. A model "
                     "fitted on this alone can only learn to output a constant, and a "
                     "stratified split has a single stratum to split. As a review list of the "
                     "double-tagged images this sheet is correct; as a training set it cannot "
                     "learn the difference between a folded and a non-folded lens, because no "
                     "non-folded lens is in it."),
        ("Option B", "All 5 columns carry both positives and negatives, so each column has "
                     "something to learn and the stratified split has real strata. The 178 "
                     "double-tagged images from Option A are all inside Option B, flagged in "
                     "the n_lp_tags column, so Option A is one filter away: "
                     "n_lp_tags = 2."),
        ("Recommendation", "Train on Option B and keep n_lp_tags in the file, so the 178 can "
                           "be isolated for review or for a dedicated evaluation at any time "
                           "without rebuilding anything."),
        ("Also in the file", "also_has_other_tag = 1 marks rows that carry tags outside the LP "
                             "family. Those tags are not modelled by this classifier; the "
                             "column only records that they exist."),
        ("Excluded for now", "Images carrying no LP tag are not in either sheet. They will be "
                             "added when the completed annotation export arrives."),
        ("Next step", "Pick A or B. The dataframe and split code then go for review before "
                      "any training run."),
    ]
    for head, body in notes:
        ws.cell(r, 1, head).font = BOLD
        ws.cell(r, 1).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(r, 2, body)
        c.alignment = WRAP
        ws.row_dimensions[r].height = max(30, 13 * (len(body) // 105 + 1))
        r += 1

    return ws


def counts_sheet(wb, a, b):
    ws = wb.create_sheet("Column counts")
    ws.append(["Label column",
               "Option A positives", "Option A negatives",
               "Option B positives", "Option B negatives"])
    for i, name in enumerate(LABEL_COLS):
        pa = sum(x["labels"][i] for x in a)
        pb = sum(x["labels"][i] for x in b)
        ws.append([name, pa, len(a) - pa, pb, len(b) - pb])
    ws.append([])
    ws.append(["Total rows", len(a), "", len(b), ""])
    style_header(ws, 5)
    set_widths(ws, [30, 20, 20, 20, 20])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=5):
        for cell in row:
            cell.border = BOX
    return ws


def patterns_sheet(wb, rows):
    ws = wb.create_sheet("Label patterns in the data")
    ws.append(["Pattern"] + LABEL_COLS + ["Rows"])
    for pat, n in Counter(tuple(x["labels"]) for x in rows).most_common():
        ws.append([" + ".join(LABEL_COLS[i] for i, v in enumerate(pat) if v)] + list(pat) + [n])
    ws.append([])
    ws.append(["Total rows", "", "", "", "", "", len(rows)])
    style_header(ws, 7)
    set_widths(ws, [46, 18, 13, 17, 25, 13, 10])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=7):
        for cell in row:
            cell.border = BOX
    return ws


def main():
    rows = load()
    a = [x for x in rows if x["n_lp_tags"] == 2]
    b = rows
    assert len(a) == 178, len(a)
    assert len(b) == 3377, len(b)

    # check the 67 / 111 / 178 split quoted from the lead
    pure = sum(1 for x in a if x["also_has_other_tag"] == 0)
    mixed = sum(1 for x in a if x["also_has_other_tag"] == 1)
    print("double-tagged rows: %d total | LP tags only: %d | LP plus other tags: %d"
          % (len(a), pure, mixed))
    print("lead quoted: 178 total | 67 pure | 111 plus other  -> %s"
          % ("MATCH" if (len(a), pure, mixed) == (178, 67, 111) else "DIFFERS"))

    write_csv(CSV_A, a)
    write_csv(CSV_B, b)
    print("wrote %s (%d rows)" % (CSV_A.name, len(a)))
    print("wrote %s (%d rows)" % (CSV_B.name, len(b)))

    wb = Workbook()
    decision_sheet(wb, a, b)
    data_sheet(wb, "Option A - 178 rows", a, "FFC000")
    data_sheet(wb, "Option B - 3377 rows", b, "70AD47")
    counts_sheet(wb, a, b)
    patterns_sheet(wb, b)
    wb.save(XLSX)
    print("wrote %s (%d sheets)" % (XLSX.name, len(wb.sheetnames)))
    print("sheets: %s" % " | ".join(wb.sheetnames))
    return 0


if __name__ == "__main__":
    sys.exit(main())

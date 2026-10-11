#!/usr/bin/env python3
"""Workbook for the Task 1 record (wayfinder ticket 006).

Sheets:
    Run summary        the four runs, headline metrics, side by side
    Bubble family      the 185 / 55 picture comparison, both architectures
    Attribution        confusion cells that name the responsible class
    Per class          every class x run: support, precision, recall, f1
    Count table        the client's count table (separate + merged, LP/OOI)

Inputs: the mirrored run artifacts under <runs> and the verified count tables
under <counts>. Needs openpyxl (repo .venv).

    .venv/Scripts/python.exe make_readout_workbook.py --runs <dir> \\
        --counts <run/bubble_merge/20261006> --out <xlsx path>
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from readout import FAMILY, RUNS, family_stats, leaks_into, prf, top_confusions  # noqa: E402

from openpyxl import Workbook                      # noqa: E402
from openpyxl.styles import Alignment, Font        # noqa: E402

BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = BOLD
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    return ws


def load(runs_dir, run_id):
    d = os.path.join(runs_dir, run_id)
    out = {"id": run_id}
    with open(os.path.join(d, "metrics.json"), encoding="utf-8") as fh:
        out["metrics"] = json.load(fh)
    with open(os.path.join(d, "confusion.json"), encoding="utf-8") as fh:
        out["confusion"] = json.load(fh)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--runs", required=True)
    ap.add_argument("--counts", required=True,
                    help="dir with count_table_separate.csv / _merged.csv")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    runs = {rid: load(args.runs, rid) for rid, _, _ in RUNS}
    wb = Workbook()
    wb.remove(wb.active)

    # --- Run summary ----------------------------------------------------
    rows = []
    for rid, arch, variant in RUNS:
        m = runs[rid]["metrics"]
        rows.append([rid, arch, variant, m["num_classes"],
                     m["num_train_images"], m["num_test_images"],
                     round(m["overall_accuracy"], 4),
                     round(m["macro"]["f1"], 4),
                     round(m["macro"]["precision"], 4),
                     round(m["macro"]["recall"], 4)])
    sheet(wb, "Run summary",
          ["run id", "architecture", "reading", "classes", "train", "test",
           "overall accuracy", "macro F1", "macro precision",
           "macro recall"],
          rows, [36, 14, 11, 8, 8, 8, 16, 11, 16, 14])

    # --- Bubble family + attribution ------------------------------------
    fam_rows, attr_rows = [], []
    for arch, sid, mid in (("ResNet50", "bubble_merge_resnet50_separate",
                            "bubble_merge_resnet50_merged"),
                           ("MultiScale", "bubble_merge_multiscale_separate",
                            "bubble_merge_multiscale_merged")):
        sep, mrg = runs[sid], runs[mid]
        fs = family_stats(sep, FAMILY)
        fm = family_stats(mrg, ["Bubble"])
        fam_rows.append([arch, "separate (4 classes)", fs["support"],
                         fs["exact_correct"], round(fs["recall"], 4),
                         round(fs["family_recall"], 4),
                         round(fs["precision"], 4)])
        fam_rows.append([arch, "merged (1 class)", fm["support"],
                         fm["exact_correct"], round(fm["recall"], 4),
                         round(fm["family_recall"], 4),
                         round(fm["precision"], 4)])
        for n, cls in top_confusions(mrg, "Bubble", limit=8):
            attr_rows.append([arch, "merged", "merged Bubble lost to", cls, n])
        for cls, n in leaks_into(mrg, "Bubble", limit=8):
            attr_rows.append([arch, "merged", "swallowed by merged Bubble",
                              cls, n])
        for b in FAMILY:
            for n, cls in top_confusions(sep, b, limit=4):
                attr_rows.append([arch, "separate", "row %s leaks to" % b,
                                  cls, n])
    sheet(wb, "Bubble family",
          ["architecture", "reading", "support (test)", "exact correct",
           "exact recall", "family recall", "precision"],
          fam_rows, [14, 22, 16, 14, 13, 14, 10])
    sheet(wb, "Attribution",
          ["architecture", "reading", "direction", "class", "pictures"],
          attr_rows, [14, 11, 28, 34, 10])

    # --- Per class ------------------------------------------------------
    per_rows = []
    for rid, arch, variant in RUNS:
        cm = runs[rid]["confusion"]
        for i, c in enumerate(cm["class_names"]):
            p, r, f, s = prf(cm["matrix"], i)
            per_rows.append([arch, variant, c, s, round(p, 4), round(r, 4),
                             round(f, 4)])
    sheet(wb, "Per class",
          ["architecture", "reading", "class", "support", "precision",
           "recall", "f1"],
          per_rows, [14, 11, 40, 10, 10, 10, 10])

    # --- Count table ----------------------------------------------------
    ct_rows = []
    for name in ("count_table_separate.csv", "count_table_merged.csv"):
        with open(os.path.join(args.counts, name), newline="",
                  encoding="utf-8") as fh:
            for i, r in enumerate(csv.DictReader(fh)):
                ct_rows.append([name.replace("count_table_", "")
                                .replace(".csv", ""), i + 1] +
                               [r[k] for k in r])
    header = ["variant", "#"] + list(r.keys())
    sheet(wb, "Count table", header, ct_rows,
          [11, 5] + [40] + [14] * (len(header) - 3))

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)
    print("wrote", args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

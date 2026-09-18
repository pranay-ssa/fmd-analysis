"""Append Corrected_OOI_20260908 and Corrected_Lens_20260908 sheets to Comparision.xlsx,
mirroring the Sheet1 column layout (Architecture / Train / Test / Total / ms-img)."""
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = "D:/02-SSA/fmd-analysis/Comparision.xlsx"

ooi = [
    ("ResNet18", 46.39, 0.0847),
    ("ResNet50", 115.61, 0.1212),
    ("ResNet50-Inception", 138.85, 0.1530),
    ("ResNet50-SE", 127.88, 0.1517),
]
lens = [
    ("ResNet18", 95.80, 0.1451),
    ("ResNet50", 244.10, 0.2759),
    ("ResNet50-Inception", 290.80, 0.3406),
    ("ResNet50-SE", 267.84, 0.3350),
]

wb = openpyxl.load_workbook(SRC)
bold = Font(bold=True)
hfill = PatternFill("solid", fgColor="D9E1F2")


def add_sheet(title, data, n_div, header):
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = bold
        c.fill = hfill
    for a, tr, te in data:
        ws.append([a, tr, te, round(tr + te, 4), round(te / n_div * 1000, 3)])
    for i, w in enumerate([20, 13, 11, 12, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    print("added", title, "rows", ws.max_row - 1)


add_sheet("Corrected_OOI_20260908", ooi, 87,
          ["Architecture", "OOI Train (s)", "OOI Test (s)", "OOI Total (s)", "OOI ms/img (÷87)"])
add_sheet("Corrected_Lens_20260908", lens, 201,
          ["Architecture", "Lens Train (s)", "Lens Test (s)", "Lens Total (s)", "Lens ms/img (÷201)"])

wb.save(SRC)
print("saved", SRC, "sheets:", wb.sheetnames)
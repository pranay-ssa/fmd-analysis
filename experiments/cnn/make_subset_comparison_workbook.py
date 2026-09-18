#!/usr/bin/env python3
"""
Workbook for the six-class question: a model trained on the six classes against the full model.

Reads run/line_results_subset6_20260917/subset_vs_full.json, written by compare_subset_vs_full.py,
and writes reports/cnn/FMD_SixClasses_vs_FullModel_20260917.xlsx.

    Sheet 1: one row per line, both models on the same test pictures, with the change
    Sheet 2: per class per line, the same comparison, negative changes flagged
"""
from __future__ import annotations

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(REPO, "run", "line_results_subset6_20260917", "subset_vs_full.json")
OUT = os.path.join(REPO, "reports", "cnn", "FMD_SixClasses_vs_FullModel_20260917.xlsx")

TITLE = Font(bold=True, size=14)
SUB = Font(bold=True, size=11)
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F3864")
NOTE = Font(italic=True, size=9, color="555555")
BOLD = Font(bold=True)
DOWN = PatternFill("solid", fgColor="FCE4E4")     # where restricting the classes scored lower
UP = PatternFill("solid", fgColor="E2EFDA")       # where it scored higher
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, row, values):
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = HEAD
        cell.fill = FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def main():
    d = json.load(open(SRC, encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Six classes vs full model"
    ws["A1"] = "The six highlighted classes: a model trained on just those six against the full model"
    ws["A1"].font = TITLE
    notes = [
        "Same test pictures in both columns. A picture kept its train or test side from the one split, so the",
        "two models answered exactly the same {0} pictures of the six classes.".format(d["pooled"]["test_pictures"]),
        "The six model only ever answers with six labels. The full model answers with all of its labels, so its",
        "figures here are read out of its full confusion matrix: recall from the class row, precision from the",
        "class column of the full matrix. A foreign matter picture the full model calls Bubble counts as wrong,",
        "which a matrix trimmed to six classes would have hidden.",
        "A positive change means restricting the training to the six classes scored higher.",
        "A cell with four or fewer test pictures is not evidence. Check the n column before reading a class.",
    ]
    r = 3
    for t in notes:
        ws.cell(row=r, column=1, value=t).font = NOTE
        r += 1
    r += 1
    header(ws, r, ["Line", "Test pictures of the six", "Full model accuracy", "Six-class accuracy",
                   "Accuracy change", "Full model macro recall", "Six-class macro recall",
                   "Macro recall change", "Full model macro F1", "Six-class macro F1", "Macro F1 change"])
    head_row = r
    r += 1
    first = r
    for row in d["lines"]:
        vals = [row["line"], row["test_pictures_of_the_six"], row["full_accuracy"], row["six_accuracy"],
                row["accuracy_gap"], row["full_macro_recall"], row["six_macro_recall"],
                row["macro_recall_gap"], row["full_macro_f1"], row["six_macro_f1"], row["macro_f1_gap"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            if c in (3, 4, 6, 7, 9, 10):
                cell.number_format = "0.00%"
            if c in (5, 8, 11):
                cell.number_format = "+0.00%;-0.00%"
                cell.fill = UP if (v or 0) > 0 else (DOWN if (v or 0) < 0 else PatternFill())
        r += 1
    n = d["pooled"]["test_pictures"]
    r += 1
    for label, acc in (("full model", d["pooled"]["full_accuracy"]),
                       ("six class model", d["pooled"]["six_accuracy"])):
        ws.cell(row=r, column=1, value="All five, pooled, {0}".format(label)).font = BOLD
        ws.cell(row=r, column=2, value=n).font = BOLD
        cell = ws.cell(row=r, column=3, value=acc)
        cell.font = BOLD
        cell.number_format = "0.00%"
        r += 1
    ws.cell(row=r, column=1,
            value="Pooled accuracy weights each line by its test pictures, the same rule as the other "
                  "workbooks.").font = NOTE
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 13
    for col in "CDEFGHIJK":
        ws.column_dimensions[col].width = 15
    ws.freeze_panes = ws.cell(row=first, column=1).coordinate

    ws2 = wb.create_sheet("Per class")
    ws2["A1"] = "The same comparison class by class"
    ws2["A1"].font = TITLE
    ws2["A3"] = ("Recall is the share of that class's own test pictures the model answered correctly. "
                 "Red means restricting the classes scored lower on that class, green higher.")
    ws2["A3"].font = NOTE
    header(ws2, 5, ["Line", "Class", "Test pictures", "Full model recall", "Six-class recall",
                    "Recall change"])
    r = 6
    for row in d["lines"]:
        for cls, v in row["per_class"].items():
            fr, sr = v.get("recall_full"), v.get("recall_six")
            change = (sr - fr) if (fr is not None and sr is not None) else None
            for c, val in enumerate([row["line"], cls, v["support"], fr, sr, change], start=1):
                cell = ws2.cell(row=r, column=c, value=val)
                cell.border = BOX
                if c in (4, 5):
                    cell.number_format = "0.00%"
                if c == 6 and change is not None:
                    cell.number_format = "+0.00%;-0.00%"
                    cell.fill = UP if change > 0 else (DOWN if change < 0 else PatternFill())
            r += 1
    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 28
    for col in "CDEF":
        ws2.column_dimensions[col].width = 16
    ws2.freeze_panes = "A6"

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote {0}".format(os.path.relpath(OUT, REPO).replace("\\", "/")))
    print("sheets: {0}".format(wb.sheetnames))
    return 0


if __name__ == "__main__":
    sys.exit(main())

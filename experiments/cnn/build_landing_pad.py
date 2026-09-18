#!/usr/bin/env python3
"""Build the landing-pad workbook matching the lead's ContactLensDefectResults.xlsx
layout, adapted for our Lens Presentation 12-class run with an Epochs column.

Column map (rows 1-2 are a two-level header):
  A            row index
  B            Dataset                (merge B1:B2)
  C            Architecture           (merge C1:C2)
  D            Overall Accuracy (%)   (merge D1:D2)
  E            Epochs                 (merge E1:E2)
  F..Q         Recall (%)             (merge F1:Q1; row2 = 12 class codes)
  R            Training time (s)      (merge R1:R2)
  S            Inference Time per Image(ms)  (merge S1:S2)

One data row per (architecture, epoch) -> 12 rows (4 arch x epochs 20/30/50).

With --metrics DIR, each row is auto-filled from
  DIR/<arch_key>/epochs_<ep>/metrics.json
(arch_key = lowercase model_training --arch value).
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LP_CLASSES = [
    "Bubble Scatter",
    "Cavity Off Center",
    "Dirty Camera",
    "Dirty Strobe",
    "HEMA Obstruction",
    "Lens Off Center",
    "Low Dose Obstructing Region of Interest",
    "Missing Lens",
    "Missing Primary Package",
    "Multiple Lenses",
    "Package Misalignment",
    "View Obstructed",
]
# Short codes used for the per-class Recall columns (like the lead's CT/D/EC...).
CLASS_CODES = [
    "BS", "COC", "DC", "DS", "HO", "LOC",
    "LDO", "ML", "MPP", "MUL", "PM", "VO",
]
EPOCHS = [20, 30, 50]
ARCHITECTURES = [
    "ResNet18",
    "ResNet50",
    "ResNet50-Inception  (Multi-scale) *",
    "ResNet50-SE  (Channel attention) **",
]
# display arch -> metrics.json --arch key
ARCH_KEYS = {
    "ResNet18": "resnet18",
    "ResNet50": "resnet50",
    "ResNet50-Inception  (Multi-scale) *": "resnet50_inception",
    "ResNet50-SE  (Channel attention) **": "resnet50_se",
}
DATASET_LABEL = ("Datasetver (LP 12-class, 20260903 fixed crops, "
                 "batch 8, train=multi / test=single)")

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def fill_row(ws, row, arch_disp, ep, metrics_dir: Path | None):
    ws.cell(row, 1, row - 2)            # row index
    ws.cell(row, 3, arch_disp)          # Architecture
    ws.cell(row, 5, ep)                 # Epochs
    if metrics_dir is None:
        return
    key = ARCH_KEYS[arch_disp]
    mpath = metrics_dir / key / f"epochs_{ep}" / "metrics.json"
    if not mpath.exists():
        ws.cell(row, 4, f"missing:{key}/ep{ep}")
        return
    d = json.loads(mpath.read_text())
    ws.cell(row, 4, round(float(d["overall_accuracy"]) * 100, 2))       # D acc%
    # per-class recall into F.. (row2 class codes aligned with LP_CLASSES)
    recall_first = 6
    for i, cls in enumerate(LP_CLASSES):
        pc = d["per_class"].get(cls)
        val = round(float(pc["recall"]) * 100, 2) if pc else ""
        ws.cell(row, recall_first + i, val)
    train_t = recall_first + len(LP_CLASSES)      # R
    infer_t = train_t + 1                          # S
    ws.cell(row, train_t, round(float(d["training_time_sec"]), 2))
    ws.cell(row, infer_t, round(float(d["inference_per_image_ms"]), 2))


def build(out: Path, metrics_dir: Path | None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    n_cls = len(LP_CLASSES)
    n_epochs = len(EPOCHS)
    n_rows = len(ARCHITECTURES) * n_epochs

    recall_first = 6   # col F
    recall_last = recall_first + n_cls - 1  # col Q
    train_t = recall_last + 1   # R
    infer_t = train_t + 1       # S

    # ---- Header (rows 1-2) ----------------------------------------------
    ws["B1"] = "Dataset"
    ws["C1"] = "Architecture"
    ws["D1"] = "Overall Accuracy (%)"
    ws["E1"] = "Epochs"
    ws.merge_cells("B1:B2"); ws.merge_cells("C1:C2")
    ws.merge_cells("D1:D2"); ws.merge_cells("E1:E2")

    ws.cell(1, recall_first, "Recall (%)")
    ws.merge_cells(start_row=1, start_column=recall_first,
                   end_row=1, end_column=recall_last)

    ws.cell(1, train_t, "Training time (s)")
    ws.cell(1, infer_t, "Inference Time per Image(ms)")
    ws.merge_cells(start_row=1, start_column=train_t,
                   end_row=2, end_column=train_t)
    ws.merge_cells(start_row=1, start_column=infer_t,
                   end_row=2, end_column=infer_t)

    for i, code in enumerate(CLASS_CODES):
        ws.cell(2, recall_first + i, code)

    # style header cells A..S rows1-2
    for r in (1, 2):
        for c in range(1, infer_t + 1):
            cell = ws.cell(r, c)
            cell.font = Font(bold=True)
            cell.alignment = CENTER
            cell.border = BORDER
            if cell.value is not None:
                cell.fill = HEADER_FILL

    # ---- Data rows: one per arch x epoch ---------------------------------
    r = 3
    for arch in ARCHITECTURES:
        for ep in EPOCHS:
            fill_row(ws, r, arch, ep, metrics_dir)
            r += 1

    # merge Dataset across each arch's 3 epoch rows
    last_data = r - 1
    for arch_i in range(len(ARCHITECTURES)):
        top = 3 + arch_i * n_epochs
        bot = top + n_epochs - 1
        ws.cell(top, 2, DATASET_LABEL)
        ws.merge_cells(start_row=top, start_column=2, end_row=bot, end_column=2)

    # borders for data region A..S
    for rr in range(3, last_data + 1):
        for cc in range(1, infer_t + 1):
            cell = ws.cell(rr, cc)
            cell.border = BORDER
            cell.alignment = CENTER

    # ---- Footnotes --------------------------------------------------------
    fn_row = last_data + 2
    ws.cell(fn_row, 3, "* Multi-scale feature extraction architecture")
    ws.cell(fn_row + 1, 3, "** Channel-based attention architecture")

    # ---- Column widths ----------------------------------------------------
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 8
    for c in range(recall_first, recall_last + 1):
        ws.column_dimensions[get_column_letter(c)].width = 7
    ws.column_dimensions[get_column_letter(train_t)].width = 14
    ws.column_dimensions[get_column_letter(infer_t)].width = 16

    wb.save(out)
    mode = "AUTO-FILLED" if metrics_dir else "BLANK"
    print(f"Wrote {out}  [{mode}]  ({len(ARCHITECTURES)} arch x {n_epochs} epochs = {n_rows} rows)")
    print("Class codes:", " ".join(f"{c}={n}" for c, n in zip(CLASS_CODES, LP_CLASSES)))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="samples/ContactLensDefectResults_ours.xlsx", type=Path)
    ap.add_argument("--metrics", type=Path, default=None,
                    help="dir containing <arch_key>/epochs_<ep>/metrics.json to auto-fill")
    args = ap.parse_args()
    build(args.out, args.metrics)

#!/usr/bin/env python3
"""Build the single-scale (train+test = single-size crops), 50-epoch-only
landing-pad workbook in the same format as ContactLensDefectResults_ours.xlsx.

One row per architecture (4 rows), epochs = 50 for all. Auto-fills from a dir
of flat files  <arch_key>.json  (each = one run's metrics.json).

USAGE
    python3 scripts/build_single50_landing_pad.py \
        --out  samples/ContactLensDefectResults_single50.xlsx \
        --metrics run/single50_metrics
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from build_landing_pad import (ARCHITECTURES, ARCH_KEYS, CLASS_CODES,
                               CENTER, BORDER, HEADER_FILL, LP_CLASSES)

EPOCH = 50
DATASET_LABEL = ("Datasetver (LP 12-class single-scale, 20260903 fixed crops, "
                 "train+test = single-size, batch 8)")


def build(out: Path, metrics_dir: Path | None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    n_cls = len(LP_CLASSES)
    recall_first = 6
    recall_last = recall_first + n_cls - 1
    train_t = recall_last + 1
    infer_t = train_t + 1

    # ---- Header (rows 1-2) ----------------------------------------------
    ws["B1"] = "Dataset"
    ws["C1"] = "Architecture"
    ws["D1"] = "Overall Accuracy (%)"
    ws["E1"] = "Epochs"
    ws.merge_cells("B1:B2"); ws.merge_cells("C1:C2")
    ws.merge_cells("D1:D2"); ws.merge_cells("E1:E2")
    ws.cell(1, recall_first, "Recall (%)")
    ws.merge_cells(start_row=1, start_column=recall_first,
                   end_row=1, end_column=recall_last)
    ws.cell(1, train_t, "Training time (s)")
    ws.cell(1, infer_t, "Inference Time per Image(ms)")
    ws.merge_cells(start_row=1, start_column=train_t, end_row=2, end_column=train_t)
    ws.merge_cells(start_row=1, start_column=infer_t, end_row=2, end_column=infer_t)
    for i, code in enumerate(CLASS_CODES):
        ws.cell(2, recall_first + i, code)

    for r in (1, 2):
        for c in range(1, infer_t + 1):
            cell = ws.cell(r, c)
            cell.font = Font(bold=True)
            cell.alignment = CENTER
            cell.border = BORDER
            if cell.value is not None:
                cell.fill = HEADER_FILL

    # ---- Data rows: one per arch (all epoch 50) --------------------------
    row = 3
    for arch in ARCHITECTURES:
        ws.cell(row, 1, row - 2)
        ws.cell(row, 3, arch)
        ws.cell(row, 5, EPOCH)
        if metrics_dir is not None:
            key = ARCH_KEYS[arch]
            f = metrics_dir / f"{key}.json"
            if f.exists():
                d = json.loads(f.read_text())
                ws.cell(row, 4, round(float(d["overall_accuracy"]) * 100, 2))
                for i, cls in enumerate(LP_CLASSES):
                    pc = d["per_class"].get(cls)
                    ws.cell(row, recall_first + i,
                            round(float(pc["recall"]) * 100, 2) if pc else "")
                ws.cell(row, train_t, round(float(d["training_time_sec"]), 2))
                ws.cell(row, infer_t, round(float(d["inference_per_image_ms"]), 2))
            else:
                ws.cell(row, 4, f"missing:{key}.json")
        row += 1

    last_data = row - 1
    ws.cell(3, 2, DATASET_LABEL)
    ws.merge_cells(start_row=3, start_column=2, end_row=last_data, end_column=2)

    for rr in range(3, last_data + 1):
        for cc in range(1, infer_t + 1):
            cell = ws.cell(rr, cc)
            cell.border = BORDER
            cell.alignment = CENTER

    fn_row = last_data + 2
    ws.cell(fn_row, 3, "* Multi-scale feature extraction architecture")
    ws.cell(fn_row + 1, 3, "** Channel-based attention architecture")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 8
    for c in range(recall_first, recall_last + 1):
        ws.column_dimensions[get_column_letter(c)].width = 7
    ws.column_dimensions[get_column_letter(train_t)].width = 14
    ws.column_dimensions[get_column_letter(infer_t)].width = 16

    wb.save(out)
    mode = "AUTO-FILLED" if metrics_dir else "BLANK"
    print(f"Wrote {out}  [{mode}]  ({len(ARCHITECTURES)} arch x epoch {EPOCH})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="samples/ContactLensDefectResults_single50.xlsx", type=Path)
    ap.add_argument("--metrics", type=Path, default=None)
    args = ap.parse_args()
    build(args.out, args.metrics)

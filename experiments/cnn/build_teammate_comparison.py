#!/usr/bin/env python3
"""Add a "Comparison vs teammate" sheet to the combined-22 workbook.

Their run: the same six architectures on a different split of the same 22-class set, no
warmup (numbers supplied 2026-09-15, stored in teammate_split_20260915.json).

What is compared: overall accuracy and all 22 per-class recalls, theirs against ours, with
the delta in percentage points (ours minus theirs). Training time is shown as a control,
since it depends on architecture, epochs and train-set size and so should be broadly equal
across the two splits.

Inference per image is deliberately EXCLUDED. Their run had no warmup, so their ms/img
carries a one-time process cost (XLA re-trace plus autotune) that ours does not, and the
two are not like-for-like. Measuring it on our own split is parked in
experiments/cnn/README.md.

USAGE
    python3 experiments/cnn/build_teammate_comparison.py \
        --workbook reports/cnn/ContactLensDefectResults_combined22.xlsx \
        --metrics  runs/cnn/combined22
"""
import argparse
import json
import pathlib

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from build_combined22_landing_pad import ALL_22_CLASSES, CLASS_CODES
from build_landing_pad import BORDER, CENTER

ARCHES = [
    ("ResNet18", "resnet18"),
    ("ResNet50", "resnet50"),
    ("ResNet50-Inception  (Multi-scale) *", "resnet50_inception"),
    ("ResNet50-SE  (Channel attention) **", "resnet50_se"),
    ("ResNet50-Inception-SE  (Branch attention) ***",
     "resnet50_inception_attention"),
    ("ResNet50-Inception-Refine  (Bicubic upsample) ****",
     "resnet50_inception_refine"),
]

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
DELTA_FILL = PatternFill("solid", fgColor="FFF2CC")
SHEET_NAME = "Comparison vs teammate"
TEAMMATE_JSON = pathlib.Path(__file__).with_name("teammate_split_20260915.json")


def build(workbook: pathlib.Path, metrics_dir: pathlib.Path) -> dict:
    tm = json.loads(TEAMMATE_JSON.read_text())
    theirs = {r["key"]: r for r in tm["rows"]}
    ours = {k: json.loads((metrics_dir / k / "epochs_50" / "metrics.json").read_text())
            for _, k in ARCHES}

    wb = load_workbook(workbook)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)

    # ---- header: A architecture | B-D accuracy | E-F training | per class 3 cols ----
    ws["A1"] = "Architecture"
    ws.merge_cells("A1:A2")
    ws["B1"] = "Overall Accuracy (%)"
    ws.merge_cells("B1:D1")
    ws["E1"] = "Training time (s)"
    ws.merge_cells("E1:F1")
    for i, code in enumerate(CLASS_CODES):
        first = 7 + i * 3
        ws.cell(1, first, f"{code} recall (%)")
        ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=first + 2)
    for col, label in ((2, "theirs"), (3, "ours"), (4, "delta"),
                       (5, "theirs"), (6, "ours")):
        ws.cell(2, col, label)
    for i in range(len(CLASS_CODES)):
        first = 7 + i * 3
        for off, label in enumerate(("theirs", "ours", "delta")):
            ws.cell(2, first + off, label)

    last_col = 6 + len(CLASS_CODES) * 3
    for r in (1, 2):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.font = Font(bold=True)
            cell.alignment = CENTER
            cell.border = BORDER
            if cell.value is not None:
                cell.fill = HEADER_FILL

    # ---- data rows ----
    row = 3
    acc_deltas = {}
    for name, key in ARCHES:
        t, o = theirs[key], ours[key]
        ws.cell(row, 1, name)
        ws.cell(row, 2, t["overall_accuracy"])
        ours_acc = round(float(o["overall_accuracy"]) * 100, 2)
        ws.cell(row, 3, ours_acc)
        ws.cell(row, 4, round(ours_acc - t["overall_accuracy"], 2))
        ws.cell(row, 5, t["training_time_sec"])
        ws.cell(row, 6, round(float(o["training_time_sec"]), 2))
        acc_deltas[key] = ours_acc - t["overall_accuracy"]
        for i, cls in enumerate(ALL_22_CLASSES):
            pc = o["per_class"].get(cls)
            ours_rec = round(float(pc["recall"]) * 100, 2) if pc else None
            their_rec = t["per_class_recall_pct"][i]
            first = 7 + i * 3
            ws.cell(row, first, their_rec)
            ws.cell(row, first + 1, ours_rec)
            if ours_rec is not None:
                ws.cell(row, first + 2, round(ours_rec - their_rec, 2))
            for off in range(3):
                ws.cell(row, first + off).fill = DELTA_FILL
        row += 1

    last_data = row - 1
    for rr in range(3, last_data + 1):
        for cc in range(1, last_col + 1):
            ws.cell(rr, cc).border = BORDER
            ws.cell(rr, cc).alignment = CENTER

    fn = last_data + 2
    notes = [
        "Comparison of the same six architectures on two different splits of the same 22-class set.",
        "theirs = teammate's split, no warmup, supplied 2026-09-15. ours = our split, warm protocol (2 dummy + 1 real-pipeline discard pass before timing).",
        "delta = ours minus theirs, in percentage points. Green-free cells (amber) are the per-class blocks.",
        "Accuracy is split-dependent, so a delta is not a quality verdict on either split; what it shows is which architectures hold their position across splits.",
        "Inference per image is excluded on purpose: their run had no warmup, so their ms/img carries a one-time process cost and is not comparable to our warm number.",
    ]
    for i, n in enumerate(notes):
        ws.cell(fn + i, 1, n)

    ws.column_dimensions["A"].width = 44
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9
    ws.freeze_panes = "B3"

    wb.save(workbook)
    return acc_deltas


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", type=pathlib.Path,
                    default=pathlib.Path("reports/cnn/ContactLensDefectResults_combined22.xlsx"))
    ap.add_argument("--metrics", type=pathlib.Path,
                    default=pathlib.Path("runs/cnn/combined22"))
    a = ap.parse_args()
    d = build(a.workbook, a.metrics)
    print(f"Added sheet '{SHEET_NAME}' to {a.workbook}")
    for k, v in d.items():
        print(f"  accuracy delta (ours - theirs) {k:<32} {v:+.2f} pp")

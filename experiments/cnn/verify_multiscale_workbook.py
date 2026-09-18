#!/usr/bin/env python3
"""Verify the multi-scale workbook against the run's own metrics.

The Summary sheet holds fractions (0.8939), not percentages, and the pooled row is
labelled "All five" in column A. This checks, for every line, that the accuracy and
macro recall in the workbook are the numbers the run itself recorded, that the pooled
row equals the pooled value recomputed from the test counts, and that the comparison
sheet carries both architectures on the same pictures.

Usage:  python experiments/cnn/verify_multiscale_workbook.py
"""
import json
import os
import sys

import openpyxl

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
XLSX = os.path.join(REPO, "reports", "cnn", "FMD_Line_MultiScale_ResNet50_20260917.xlsx")
RUNS = os.path.join(REPO, "runs", "cnn", "lines_multiscale")
BASE = os.path.join(REPO, "runs", "cnn", "lines")
LINES = ["L24", "L25", "L26", "L27", "L31"]
STANDARD_SHEETS = {"Summary", "Cells by support", "Line summary (CNN format)"}

checks = 0
fails = []


def check(label, ok, detail=""):
    global checks
    checks += 1
    if not ok:
        fails.append("{0} {1}".format(label, detail))


def numbers(ws, row, digits=4):
    out = set()
    for c in ws[row]:
        if isinstance(c.value, (int, float)):
            out.add(round(float(c.value), digits))
    return out


def metrics(root, line):
    return json.load(open(os.path.join(root, line, "epochs_50", "metrics.json")))


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    print("sheets:", wb.sheetnames)
    for name in ["Summary", "Cells by support", "Line summary (CNN format)"]:
        check("sheet present: " + name, name in wb.sheetnames)

    ws = wb["Summary"]
    hdr = None
    for r in range(1, 40):
        vals = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[r]]
        if "line" in vals and "train" in vals:
            hdr = r
            break
    check("summary header row found", hdr is not None)
    if hdr is None:
        return report()

    rows = {}
    for r in range(hdr + 1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if isinstance(label, str) and label.strip() in LINES and label.strip() not in rows:
            rows[label.strip()] = r
    check("all five lines listed", len(rows) == 5, str(sorted(rows)))

    pooled_acc = pooled_n = 0.0
    total_train = 0
    for line in LINES:
        m = metrics(RUNS, line)
        total_train += m["num_train_images"]
        pooled_acc += m["overall_accuracy"] * m["num_test_images"]
        pooled_n += m["num_test_images"]
        r = rows.get(line)
        if r is None:
            check(line + " row present", False)
            continue
        nums = numbers(ws, r)
        check("{0} accuracy {1}".format(line, round(m["overall_accuracy"], 4)),
              round(m["overall_accuracy"], 4) in nums, "row " + str(r))
        check("{0} macro recall {1}".format(line, round(m["macro"]["recall"], 4)),
              round(m["macro"]["recall"], 4) in nums, "row " + str(r))
        check("{0} test count {1}".format(line, m["num_test_images"]),
              m["num_test_images"] in nums, "row " + str(r))

    pooled = round(pooled_acc / pooled_n, 4)
    pooled_row = None
    for r in range(hdr + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and "all five" in v.lower():
            pooled_row = r
            break
    check("pooled row found", pooled_row is not None)
    if pooled_row:
        check("pooled accuracy {0}".format(pooled), pooled in numbers(ws, pooled_row),
              "row " + str(pooled_row))

    text = " ".join(str(c.value) for row in ws.iter_rows(min_row=1, max_row=hdr)
                    for c in row if c.value is not None)
    check("protocol names this architecture",
          "MultiLevel" in text or "multi-scale" in text.lower())
    check("protocol states total train {0:,}".format(total_train),
          "{:,}".format(total_train) in text)
    check("protocol states pooled test count", "{:,}".format(int(pooled_n)) in text)

    # the comparison sheet, named by the builder and truncated by Excel's 31-character cap
    extra = [s for s in wb.sheetnames if s not in STANDARD_SHEETS and s not in LINES]
    check("comparison sheet exists", len(extra) >= 1, str(extra))
    if extra:
        wsc = wb[extra[0]]
        for line in LINES:
            mb, mm = metrics(BASE, line), metrics(RUNS, line)
            pair = {round(mb["overall_accuracy"], 4), round(mm["overall_accuracy"], 4)}
            hit = False
            for r in range(1, wsc.max_row + 1):
                if pair <= numbers(wsc, r):
                    hit = True
                    break
            check("comparison sheet holds both architectures for " + line, hit)

    return report()


def report():
    print()
    if fails:
        print("FAIL {0}/{1} checks".format(len(fails), checks))
        for f in fails:
            print("  -", f)
        return 1
    print("PASS {0}/{0} checks".format(checks))
    return 0


if __name__ == "__main__":
    sys.exit(main())
